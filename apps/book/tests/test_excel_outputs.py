# pylint: disable=E0401,W0621
"""FR : Tests des exports Excel de l'app book


EN : Tests for book Excel exports

created at: 2026-10-11
created by: Paulo ALVES
"""

import io

import pytest
from openpyxl import load_workbook

from apps.book.excel_outputs import (
    book_excel_errors_emails,
    book_excel_societies_list,
    book_excel_statistiques_list,
    book_excel_supplier_cct,
)
from apps.book.excel_outputs.book_columns import SocietiesColumns, columns_list_statistiques
from apps.book.imports.excel_generations import writre_book_files
from apps.book.models import Society

pytestmark = pytest.mark.django_db


def _sheet_rows(file_io):
    file_io.seek(0)
    sheet = load_workbook(file_io, read_only=True).active
    return [tuple(cell for cell in row) for row in sheet.iter_rows(values_only=True)]


@pytest.fixture
def societies(make_society, make_payment_condition, make_country):
    make_country("FR")
    make_society(
        "FOUR001",
        name="Fournisseur",
        is_client=False,
        is_supplier=True,
        is_transporter=True,
        payment_condition_supplier=make_payment_condition("30J", name="30 jours"),
        email_01="four@acuitis.com",
    )
    make_society(
        "CLI001",
        name="Client",
        is_client=True,
        payment_condition_client=make_payment_condition("COMPTANT", name="Comptant"),
    )
    make_society("DIV001", name="Divers", is_client=False, is_various=True)


# TIERS ============================================================================================


@pytest.mark.parametrize(
    ("society_type", "expected_tiers", "nb_columns"),
    [
        ("tiers", ["CLI001", "DIV001", "FOUR001"], 36),
        ("clients", ["CLI001"], 28),
        ("suppliers", ["FOUR001"], 28),
    ],
)
def test_lignes_des_tiers_par_type(societies, society_type, expected_tiers, nb_columns):
    """Les lignes SQL correspondent aux colonnes déclarées pour chaque type de liste"""
    rows = book_excel_societies_list.GetRows(Society, society_type).get_clean_rows()
    columns = getattr(SocietiesColumns, f"columns_list_{society_type}")

    assert sorted(row[0] for row in rows) == expected_tiers
    assert len(columns) == nb_columns
    assert all(len(row) == nb_columns for row in rows)


def test_lignes_des_tiers_contenu(societies):
    rows = {
        row[0]: row for row in book_excel_societies_list.GetRows(Society, "tiers").get_clean_rows()
    }

    assert rows["FOUR001"][16] == "30 jours"
    assert rows["FOUR001"][22:31] == ("", "", "", "X", "", "", "X", "", "")
    assert rows["FOUR001"][31] == "four@acuitis.com"
    assert rows["CLI001"][19] == "Comptant"
    assert rows["CLI001"][22] == "X"
    assert rows["DIV001"][26] == "X"
    assert book_excel_societies_list.GetRows(Society, "clients").get_clean_rows()[0][22] == "X"
    assert book_excel_societies_list.GetRows(Society, "suppliers").get_clean_rows()[0][22] == "X"


def test_generation_excel_des_tiers(societies):
    file_io = io.BytesIO()

    result = book_excel_societies_list.excel_liste_societies(
        file_io, "LISTING_DES_FOURNISSEURS.xlsx", Society, "suppliers"
    )

    rows = _sheet_rows(file_io)
    assert result == {
        "OK": "GENERATION DU FICHIER LISTING_DES_FOURNISSEURS.xlsx TERMINEE AVEC SUCCES"
    }
    assert rows[0][0] == "LISTING DES FOURNISSEURS.xlsx"
    assert rows[3][0] == "N° de tiers x3"
    assert rows[4][:2] == ("FOUR001", "Fournisseur")


def test_writre_book_files_ecrit_les_trois_listes(societies, settings, tmp_path):
    """Les trois listes volumineuses sont générées sur disque, prêtes au téléchargement"""
    settings.MEDIA_EXCEL_FILES_DIR = tmp_path

    writre_book_files()

    assert sorted(path.name for path in tmp_path.iterdir()) == [
        "LISTING_DES_CLIENTS.xlsx",
        "LISTING_DES_FOURNISSEURS.xlsx",
        "LISTING_DES_TIERS.xlsx",
    ]
    rows = _sheet_rows(io.BytesIO((tmp_path / "LISTING_DES_CLIENTS.xlsx").read_bytes()))
    assert rows[4][0] == "CLI001"
    assert len(rows) == 5


# IDENTIFIANTS CCT =================================================================================


def test_excel_identifiants_cct(make_society, make_cct, make_supplier_cct):
    society = make_society("FOUR001", name="Fournisseur", is_client=False)
    make_supplier_cct(society, make_cct("AF0002", name="Lyon"), "LYO|")
    make_supplier_cct(society, make_cct("AF0001", name="Paris"), "PAR|P2|")
    make_supplier_cct(make_society("FOUR002", is_client=False), make_cct("AF0001"), "X|")
    file_io = io.BytesIO()

    result = book_excel_supplier_cct.excel_liste_supplier_cct(
        file_io, "Identifiants_cct_FOUR001.xlsx", {"third_party_num": "FOUR001"}
    )

    rows = _sheet_rows(file_io)
    assert result == {
        "OK": "GENERATION DU FICHIER Identifiants_cct_FOUR001.xlsx TERMINEE AVEC SUCCES"
    }
    assert rows[0][0] == "Identifiants founisseur pour les CCT : FOUR001 - Fournisseur"
    assert rows[3][:3] == ("Maison", "CCT X3", "Identifiants")
    assert rows[4][:3] == ("Lyon", "AF0002", "LYO|")
    assert rows[5][:3] == ("Paris", "AF0001", "PAR|P2|")
    assert len(rows) == 6


# EMAILS EN ERREUR =================================================================================


def test_excel_emails_en_erreur(make_society):
    make_society("TIERS01", email_01="bon@acuitis.com", email_02="faux")
    make_society("TIERS02", email_05="encore faux")
    file_io = io.BytesIO()

    result = book_excel_errors_emails.excel_errors_emails(
        file_io, "LISTING_DES_EMAILS_EN_ERREUR.xlsx"
    )

    rows = _sheet_rows(file_io)
    assert result == {
        "OK": "GENERATION DU FICHIER LISTING_DES_EMAILS_EN_ERREUR.xlsx TERMINEE AVEC SUCCES"
    }
    assert book_excel_errors_emails.get_rows() == [("TIERS01", "faux"), ("TIERS02", "encore faux")]
    assert rows[0][0] == "Emails tiers en erreur dans Sage X3 (Tiers)"
    assert rows[3][:2] == ("Tiers X3", "Email en erreur")
    assert rows[4][:2] == ("TIERS01", "faux")
    assert rows[5][:2] == ("TIERS02", "encore faux")


# STATISTIQUES =====================================================================================


def test_excel_statistiques(
    make_stat_familly, make_supplier_family_axes, make_sub_category, make_unit
):
    """Familles avec leurs définitions, famille sans définition conservée (jointure externe)"""
    make_stat_familly("ACUITIS")
    make_stat_familly("VIDE", description="Sans définition", regex_bool=True)
    make_supplier_family_axes(
        "LUN",
        stat_name="ACUITIS",
        description="Lunettes",
        norme="N1",
        comment="c",
        customs_code="9003",
        sub_category=make_sub_category("Montures"),
        item_weight="1.5",
        unit_weight=make_unit(1, "U"),
    )
    file_io = io.BytesIO()

    result = book_excel_statistiques_list.excel_liste_statistiques(
        file_io, "LISTING_DES_FAMILLES_STATISTIQUES_AXES_2026_10_11_1.xlsx"
    )

    rows = list(book_excel_statistiques_list.get_row())
    sheet = _sheet_rows(file_io)
    assert result == {
        "OK": "GENERATION DU FICHIER LISTING_DES_FAMILLES_STATISTIQUES_AXES_2026_10_11_1.xlsx TERMINEE AVEC SUCCES"
    }
    assert len(rows) == 2
    assert all(len(row) == len(columns_list_statistiques) for row in rows)
    assert rows[0] == (
        "ACUITIS",
        "Statistique ACUITIS",
        "",
        "famille",
        "LUN",
        "LUN",
        "OPT",
        "Lunettes",
        "N1",
        "c",
        "Marchandises",
        "Montures",
        1.5,
        1,
        "9003",
    )
    assert rows[1] == (
        "VIDE",
        "Sans définition",
        "X",
        None,
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
    )
    assert sheet[0][0] == "LISTING DES FAMILLES STATISTIQUES AXES"
    assert sheet[3][:2] == ("Familles", None)
    assert sheet[5][0] == "ACUITIS"
