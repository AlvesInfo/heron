# pylint: disable=E0401,W0621
"""FR : Tests des vues des dictionnaires : types de facture, regroupements de facturation,


     axe pro / regroupement, axe pro / tva, comptes achat vente
EN : Tests for dictionaries views

created at: 2026-10-11
created by: Paulo ALVES
"""

import io
import logging

import pytest
from django.contrib.messages import get_messages
from django.urls import reverse
from openpyxl import load_workbook

from apps.articles.models import ArticleAccount
from apps.centers_purchasing.excel_outputs.output_excel_account_axe_list import (
    excel_liste_account_axe,
)
from apps.centers_purchasing.models import (
    AccountsAxeProCategory,
    ApplicableProVat,
    AxeProGroupingGoods,
    GroupingGoods,
    TypeFacture,
)
from apps.centers_purchasing.views import (
    account_axe_pro_views,
    axe_grouping_views,
    axe_pro_vat_views,
    grouping_goods_views,
)
from apps.core.models import ChangesTrace
from conftest import patched_attrs

pytestmark = pytest.mark.django_db


def _messages(response):
    return [message.message for message in get_messages(response.wsgi_request)]


def _url(name):
    return reverse(f"centers_purchasing:{name}")


# TYPES DE FACTURE =================================================================================


def _set_centers(client, code):
    client.user.code_child_center = code
    client.user.save()


def test_liste_des_types_de_facture_filtree_par_centrale(logged_client, make_type_facture):
    """L'utilisateur ne voit que ses centrales, * les voit toutes"""
    make_type_facture("ACF", "380")
    make_type_facture("GAF", "380")
    _set_centers(logged_client, "GAF|MAA")

    response = logged_client.get(_url("type_facture_list"))
    _set_centers(logged_client, "*")
    everything = logged_client.get(_url("type_facture_list"))

    assert response.status_code == 200
    assert response.context["titre_table"] == "Centrale / Type de Facture"
    assert [row.child_center.code for row in response.context["type_factures"]] == ["GAF"]
    assert [row.child_center.code for row in everything.context["type_factures"]] == ["ACF", "GAF"]


def test_creation_type_de_facture(logged_client, make_child_center):
    make_child_center("ACF")
    data = {
        "child_center": "ACF",
        "invoice_type": "380",
        "purchase_type_facture": "FAC",
        "sale_type_facture": "FAV",
    }

    get = logged_client.get(_url("type_facture_create"))
    response = logged_client.post(_url("type_facture_create"), data)
    invalid = logged_client.post(_url("type_facture_create"), {**data, "sale_type_facture": ""})

    assert get.context["create"] is True
    assert get.context["titre_table"] == "Création d'un nouveau type de Facture"
    assert response.status_code == 302
    assert response.url == _url("type_facture_list")
    assert TypeFacture.objects.get().purchase_type_facture == "FAC"
    assert (
        "ACF - Acuitis - 380,  a été créé avec success" in _messages(response)[0]
        or "a été créé avec success" in _messages(response)[0]
    )
    assert invalid.status_code == 200
    assert logged_client.session["level"] == 50


def test_modification_type_de_facture(logged_client, make_type_facture):
    facture = make_type_facture("ACF", "380")
    url = reverse("centers_purchasing:type_facture_update", args=[facture.pk])

    get = logged_client.get(url)
    response = logged_client.post(
        url,
        {
            "child_center": "ACF",
            "invoice_type": "380",
            "purchase_type_facture": "FAX",
            "sale_type_facture": "FAV",
        },
    )

    facture.refresh_from_db()
    assert "Mise à jour du Type de Facture : " in get.context["titre_table"]
    assert response.status_code == 302
    assert facture.purchase_type_facture == "FAX"
    assert ChangesTrace.objects.count() == 1


def test_suppression_type_de_facture(ajax_client, logged_client, make_type_facture, caplog):
    facture = make_type_facture("ACF", "380")

    with caplog.at_level(logging.ERROR, logger="views"):
        invalid = ajax_client.post(_url("type_facture_delete"), {"pk": facture.pk + 1000})
    response = ajax_client.post(_url("type_facture_delete"), {"pk": facture.pk})

    assert invalid.json() == {"success": "ko"}
    assert "type_piece_delete, form invalid" in caplog.text
    assert response.json() == {"success": "success"}
    assert TypeFacture.objects.count() == 0
    assert logged_client.post(_url("type_facture_delete"), {"pk": 1}).url == reverse("home")
    assert ajax_client.get(_url("type_facture_delete")).url == reverse("home")


# REGROUPEMENTS DE FACTURATION =====================================================================


def test_regroupements_liste_creation_modification(logged_client, make_grouping_goods):
    make_grouping_goods("Montures", ranking=1)

    liste = logged_client.get(_url("grouping_goods_list"))
    get = logged_client.get(_url("grouping_goods_create"))
    response = logged_client.post(
        _url("grouping_goods_create"),
        {"ranking": "2", "base": "Marchandises", "grouping_goods": "Verres"},
    )

    assert [row.grouping_goods for row in liste.context["grouping_goods"]] == ["Montures"]
    assert get.context["ranking"] == 2
    assert get.context["create"] is True
    assert response.status_code == 302
    assert response.url == _url("grouping_goods_list")
    created = GroupingGoods.objects.get(grouping_goods="Verres")
    assert created.created_by == logged_client.user

    url = reverse("centers_purchasing:grouping_goods_update", args=[created.pk])
    update_get = logged_client.get(url)
    update = logged_client.post(
        url, {"ranking": "3", "base": "Marchandises", "grouping_goods": "Verres"}
    )

    created.refresh_from_db()
    assert update_get.context["ranking"] == 2
    assert update_get.context["titre_table"] == "Mise à jour d'un regroupement de facturation"
    assert update.status_code == 302
    assert (created.ranking, created.modified_by) == (3, logged_client.user)


def test_regroupements_sans_ligne_rang_initial(logged_client):
    assert logged_client.get(_url("grouping_goods_create")).context["ranking"] == 1


def test_regroupements_suppression_et_export(
    ajax_client, logged_client, make_grouping_goods, monkeypatch, caplog
):
    grouping = make_grouping_goods("Montures")

    with caplog.at_level(logging.ERROR, logger="views"):
        invalid = ajax_client.post(_url("grouping_goods_delete"), {"pk": ""})
    deleted = ajax_client.post(_url("grouping_goods_delete"), {"pk": grouping.pk})
    export = logged_client.get(_url("grouping_goods_export_list"))

    assert invalid.json() == {"success": "ko"}
    assert "grouping_goods_delete, form invalid" in caplog.text
    assert deleted.json() == {"success": "success"}
    assert GroupingGoods.objects.count() == 0
    assert logged_client.post(_url("grouping_goods_delete"), {"pk": 1}).url == reverse("home")
    assert export["Content-Disposition"].startswith(
        "attachment; filename=LISTING_DES_REGROUPEMENTS_DE_FACTURATION_"
    )

    def _boom(*_args, **_kwargs):
        raise RuntimeError("excel cassé")

    with (
        patched_attrs(monkeypatch, grouping_goods_views, excel_liste_grouping_goods=_boom),
        caplog.at_level(logging.ERROR, logger="views"),
    ):
        error = logged_client.get(_url("grouping_goods_export_list"))

    assert error.url == _url("grouping_goods_list")
    assert "grouping_goods_export_list" in caplog.text


# AXE PRO / REGROUPEMENT ===========================================================================


def test_axe_regroupement_liste_creation_modification(
    logged_client, make_section, make_grouping_goods, make_axe_pro_grouping
):
    existing = make_axe_pro_grouping("OPT")
    section = make_section("PRO", "AUD")
    grouping = make_grouping_goods("Audio", ranking=2)

    liste = logged_client.get(_url("axe_grouping_list"))
    get = logged_client.get(_url("axe_grouping_create"))
    response = logged_client.post(
        _url("axe_grouping_create"),
        {
            "axe_pro": str(section.uuid_identification),
            "grouping_goods": str(grouping.uuid_identification),
        },
    )
    duplicate = logged_client.post(
        _url("axe_grouping_create"),
        {
            "axe_pro": str(existing.axe_pro.uuid_identification),
            "grouping_goods": str(grouping.uuid_identification),
        },
    )

    assert [row.axe_pro.section for row in liste.context["axes_groupings"]] == ["OPT"]
    assert get.context["titre_table"] == "Création d'un nouveau regroupement de facturation"
    assert response.status_code == 302
    assert response.url == _url("axe_grouping_list")
    assert AxeProGroupingGoods.objects.get(axe_pro=section).grouping_goods == grouping
    assert duplicate.status_code == 200
    assert "axe_pro" in duplicate.context["form"].errors

    url = reverse("centers_purchasing:axe_grouping_update", args=[existing.pk])
    update = logged_client.post(
        url,
        {
            "axe_pro": str(existing.axe_pro.uuid_identification),
            "grouping_goods": str(grouping.uuid_identification),
        },
    )

    existing.refresh_from_db()
    assert (
        logged_client.get(url).context["titre_table"]
        == "Mise à jour d'un regroupement de facturation"
    )
    assert update.status_code == 302
    assert (existing.grouping_goods, existing.modified_by) == (grouping, logged_client.user)


def test_axe_regroupement_suppression_et_export(
    ajax_client, logged_client, make_axe_pro_grouping, monkeypatch, caplog
):
    link = make_axe_pro_grouping("OPT")

    with caplog.at_level(logging.ERROR, logger="views"):
        invalid = ajax_client.post(_url("axe_grouping_delete"), {"pk": link.pk + 1000})
    deleted = ajax_client.post(_url("axe_grouping_delete"), {"pk": link.pk})
    export = logged_client.get(_url("axe_grouping_export_list"))

    assert invalid.json() == {"success": "ko"}
    assert "axe_grouping_delete, form invalid" in caplog.text
    assert deleted.json() == {"success": "success"}
    assert AxeProGroupingGoods.objects.count() == 0
    assert ajax_client.get(_url("axe_grouping_delete")).url == reverse("home")
    assert export["Content-Disposition"].startswith(
        "attachment; filename=LISTING_DES_AXE_PRO_VS_REGROUPEMENTS_DE_FACTURATION_"
    )

    def _boom(*_args, **_kwargs):
        raise RuntimeError("excel cassé")

    with (
        patched_attrs(monkeypatch, axe_grouping_views, excel_liste_axe_grouping=_boom),
        caplog.at_level(logging.ERROR, logger="views"),
    ):
        error = logged_client.get(_url("axe_grouping_export_list"))

    assert error.url == _url("axe_grouping_list")


# AXE PRO / TVA ====================================================================================


def test_axe_pro_tva_liste_creation_modification(
    logged_client, make_child_center, make_section, make_vat, make_applicable_pro_vat
):
    existing = make_applicable_pro_vat("ACF", "OPT", "001")
    make_vat("002")
    section = make_section("PRO", "AUD")

    liste = logged_client.get(_url("axe_pro_vat_list"))
    get = logged_client.get(_url("axe_pro_vat_create"))
    response = logged_client.post(
        _url("axe_pro_vat_create"),
        {"child_center": "ACF", "axe_pro": str(section.uuid_identification), "vat": "002"},
    )
    invalid = logged_client.post(
        _url("axe_pro_vat_create"), {"child_center": "ACF", "axe_pro": "", "vat": ""}
    )

    assert [row.axe_pro.section for row in liste.context["axes_pro_vat"]] == ["OPT"]
    assert (
        get.context["titre_table"]
        == "Création d'un nouveau dictionnaire Centrale Filles/Axe Pro/TVA de facturation"
    )
    assert response.status_code == 302
    assert response.url == _url("axe_pro_vat_list")
    assert ApplicableProVat.objects.get(axe_pro=section).vat.vat == "002"
    assert invalid.status_code == 200

    url = reverse("centers_purchasing:axe_pro_vat_update", args=[existing.pk])
    update = logged_client.post(
        url,
        {"child_center": "ACF", "axe_pro": str(existing.axe_pro.uuid_identification), "vat": "002"},
    )

    existing.refresh_from_db()
    assert "Mise à jour d'un dictionnaire" in logged_client.get(url).context["titre_table"]
    assert update.status_code == 302
    assert (existing.vat.vat, existing.modified_by) == ("002", logged_client.user)


def test_axe_pro_tva_suppression_et_export(
    ajax_client, logged_client, make_applicable_pro_vat, monkeypatch, caplog
):
    link = make_applicable_pro_vat("ACF", "OPT", "001")

    with caplog.at_level(logging.ERROR, logger="views"):
        invalid = ajax_client.post(_url("axe_pro_vat_delete"), {"pk": link.pk + 1000})
    deleted = ajax_client.post(_url("axe_pro_vat_delete"), {"pk": link.pk})
    export = logged_client.get(_url("axe_pro_vat_export_list"))

    assert invalid.json() == {"success": "ko"}
    assert "axe_pro_vat_delete, form invalid" in caplog.text
    assert deleted.json() == {"success": "success"}
    assert ApplicableProVat.objects.count() == 0
    assert ajax_client.get(_url("axe_pro_vat_delete")).url == reverse("home")
    assert export["Content-Disposition"].startswith(
        "attachment; filename=LISTING_DES_CENTRALE_AXE_PRO_TVA_"
    )

    def _boom(*_args, **_kwargs):
        raise RuntimeError("excel cassé")

    with (
        patched_attrs(monkeypatch, axe_pro_vat_views, excel_liste_axe_vat=_boom),
        caplog.at_level(logging.ERROR, logger="views"),
    ):
        error = logged_client.get(_url("axe_pro_vat_export_list"))

    assert error.url == _url("axe_pro_vat_list")


# COMPTES ACHAT VENTE ==============================================================================


def _account_data(account, **overrides):
    data = {
        "child_center": account.child_center.code,
        "big_category": str(account.big_category.uuid_identification),
        "sub_category": str(account.sub_category.uuid_identification)
        if account.sub_category
        else "",
        "axe_pro": str(account.axe_pro.uuid_identification),
        "vat": account.vat.vat,
        "purchase_account": str(account.purchase_account.uuid_identification),
        "sale_account": str(account.sale_account.uuid_identification),
    }
    data.update(overrides)
    return data


@pytest.fixture
def accounts(make_accounts_axe_pro_category, make_sub_category):
    """Trois paramétrages dont deux doublons ACF (rubrique nulle) et un GAF avec rubrique"""
    first = make_accounts_axe_pro_category(child_center="ACF", vat="001")
    duplicate = make_accounts_axe_pro_category(child_center="ACF", vat="001")
    gaf = make_accounts_axe_pro_category(
        child_center="GAF", vat="001", sub_category=make_sub_category("Montures")
    )
    return {"first": first, "duplicate": duplicate, "gaf": gaf}


def test_liste_des_comptes_supprime_les_doublons_et_filtre_par_centrale(logged_client, accounts):
    """L'affichage purge les doublons (le plus récent reste) ; chaque utilisateur ne voit que ses centrales"""
    _set_centers(logged_client, "GAF")

    response = logged_client.get(_url("account_axe_list"))
    _set_centers(logged_client, "*")
    everything = logged_client.get(_url("account_axe_list"))

    rows = list(everything.context["accounts_axes"])
    assert response.status_code == 200
    assert [row["child_center__code"] for row in response.context["accounts_axes"]] == ["GAF"]
    assert [row["child_center__code"] for row in rows] == ["ACF", "GAF"]
    assert rows[0]["child_center_c"] == "ACF - Centrale ACF"
    assert rows[0]["big_category"].endswith(" - Abonnements")
    assert rows[1]["sub_category"].endswith(" - Montures")
    assert rows[0]["purchase_account"] == "FRA - 607000"
    assert not AccountsAxeProCategory.objects.filter(pk=accounts["first"].pk).exists()
    assert AccountsAxeProCategory.objects.filter(pk=accounts["duplicate"].pk).exists()


def test_creation_compte(logged_client, accounts, make_vat, make_article, make_society):
    """La création met à jour les comptes des articles sans compte pour cette centrale / tva"""
    first = accounts["first"]
    make_vat("002")
    article = make_article(
        "ART-1",
        society=make_society("FOUR001", is_client=False),
        big_category=first.big_category,
        axe_pro=first.axe_pro,
    )

    get = logged_client.get(_url("account_axe_create"))
    response = logged_client.post(_url("account_axe_create"), _account_data(first, vat="002"))
    duplicate = logged_client.post(_url("account_axe_create"), _account_data(first))

    assert (
        get.context["titre_table"] == "Création d'un nouveau Dictionnaire des Comptes achat vente"
    )
    assert response.status_code == 302
    assert response.url == _url("account_axe_list")
    assert AccountsAxeProCategory.objects.filter(vat="002").exists()
    assert set(
        ArticleAccount.objects.filter(article=article).values_list("vat__vat", flat=True)
    ) == {"001", "002"}
    assert duplicate.status_code == 200
    assert duplicate.context["form"].errors["__all__"] == ["Un élément identique existe déjà!"]


def test_modification_compte(logged_client, accounts, make_account):
    first = accounts["first"]
    AccountsAxeProCategory.objects.filter(pk=accounts["duplicate"].pk).delete()
    other = make_account("FRA", "607100")
    url = reverse("centers_purchasing:account_axe_update", args=[first.pk])

    get = logged_client.get(url)
    response = logged_client.post(
        url, _account_data(first, purchase_account=str(other.uuid_identification))
    )

    first.refresh_from_db()
    assert get.context["titre_table"] == "Mise à jour d'un Dictionnaire des Comptes achat vente"
    assert response.status_code == 302
    assert (first.purchase_account, first.modified_by) == (other, logged_client.user)


def test_suppression_compte(ajax_client, logged_client, accounts, caplog):
    first = accounts["first"]

    with caplog.at_level(logging.ERROR, logger="views"):
        invalid = ajax_client.post(_url("account_axe_delete"), {"pk": first.pk + 1000})
    deleted = ajax_client.post(_url("account_axe_delete"), {"pk": first.pk})

    assert invalid.json() == {"success": "ko"}
    assert "account_axe_delete, form invalid" in caplog.text
    assert deleted.json() == {"success": "success"}
    assert not AccountsAxeProCategory.objects.filter(pk=first.pk).exists()
    assert logged_client.post(_url("account_axe_delete"), {"pk": first.pk}).url == reverse("home")


def _xlsx_first_column(content):
    sheet = load_workbook(io.BytesIO(content), read_only=True).active
    return [row[0] for row in sheet.iter_rows(min_row=5, values_only=True)]


def test_export_des_comptes_par_centrale_de_l_utilisateur(
    logged_client, accounts, monkeypatch, caplog
):
    _set_centers(logged_client, "GAF")

    response = logged_client.get(_url("account_axe_export_list"))
    _set_centers(logged_client, "*")
    everything = logged_client.get(_url("account_axe_export_list"))

    assert response["Content-Disposition"].startswith(
        "attachment; filename=LISTING_DES_AXE_PRO_VS_REGROUPEMENTS_DE_FACTURATION_"
    )
    assert _xlsx_first_column(response.content) == ["GAF"]
    assert _xlsx_first_column(everything.content) == ["ACF", "ACF", "GAF"]

    def _boom(*_args, **_kwargs):
        raise RuntimeError("excel cassé")

    with (
        patched_attrs(monkeypatch, account_axe_pro_views, excel_liste_account_axe=_boom),
        caplog.at_level(logging.ERROR, logger="views"),
    ):
        error = logged_client.get(_url("account_axe_export_list"))

    assert error.url == _url("account_axe_list")
    assert "account_axe_export_list" in caplog.text


def test_import_des_comptes_depuis_le_repertoire(logged_client, media_dirs, accounts, make_account):
    """Les classeurs déposés dans IMPORT_ACCOUNTS sont importés (upsert) puis sauvegardés"""
    first = accounts["first"]
    AccountsAxeProCategory.objects.filter(pk=accounts["duplicate"].pk).delete()
    file_io = io.BytesIO()
    excel_liste_account_axe(file_io, "LISTING_DES_COMPTES_2026_10_11_1.xlsx", "ACF")
    (media_dirs.import_accounts / "LISTING_DES_COMPTES_2026_10_11_1.xlsx").write_bytes(
        file_io.getvalue()
    )
    first.purchase_account = make_account("FRA", "607100")
    first.save()
    (media_dirs.import_accounts / "notes.txt").write_text("ignoré", encoding="utf8")

    response = logged_client.get(_url("account_axe_import_file"))

    imported = AccountsAxeProCategory.objects.get(child_center="ACF")
    assert response.status_code == 302
    assert response.url == _url("account_axe_list")
    assert imported.pk != first.pk
    assert imported.purchase_account.account == "607000"
    assert [path.name for path in media_dirs.import_accounts.iterdir()] == ["notes.txt"]
    assert (media_dirs.backup_accounts / "LISTING_DES_COMPTES_2026_10_11_1.xlsx").exists()


def test_import_des_comptes_en_erreur(logged_client, media_dirs, monkeypatch, caplog):
    def _boom(_file):
        raise RuntimeError("import cassé")

    (media_dirs.import_accounts / "comptes.xlsx").write_bytes(b"x")

    with (
        patched_attrs(monkeypatch, account_axe_pro_views, axe_pro_account=_boom),
        caplog.at_level(logging.ERROR, logger="views"),
    ):
        response = logged_client.get(_url("account_axe_import_file"))

    assert response.status_code == 302
    assert "account_axe_import_file" in caplog.text
    assert "import cassé" in caplog.text
