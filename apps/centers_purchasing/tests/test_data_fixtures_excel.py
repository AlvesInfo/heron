# pylint: disable=E0401,W0621
"""
FR : Tests des fixtures de données (centrales, enseignes) et des exports Excel de la centrale d'achat
EN : Tests for purchasing center data fixtures and Excel exports

created at: 2026-10-11
created by: Paulo ALVES
"""
import csv
import io
import logging
from pathlib import Path

import pytest
from openpyxl import load_workbook

from apps.centers_purchasing.data_fixtures.sign_board_fixtures import fixtures
from apps.centers_purchasing.excel_outputs import (
    output_excel_account_axe_list,
    output_excel_axe_grouping_list,
    output_excel_axe_pro_vat_list,
    output_excel_enseignes_list,
    output_excel_filles_list,
    output_excel_grouping_goods_list,
    output_excel_meres_list,
)
from apps.centers_purchasing.excel_outputs.output_excel_enseignes_columns import columns_list_enseignes
from apps.centers_purchasing.excel_outputs.output_excel_filles_columns import columns_list_filles
from apps.centers_purchasing.excel_outputs.output_excel_meres_columns import columns_list_meres
from apps.centers_purchasing.models import ChildCenterPurchase, PrincipalCenterPurchase, Signboard

pytestmark = pytest.mark.django_db


def _csv_codes(name):
    """Codes (première colonne) d'un fichier de fixtures de l'app"""
    path = Path(__file__).resolve().parents[1] / "data_fixtures" / name

    with path.open(encoding="utf8") as file:
        return {row[0] for row in csv.reader(file, delimiter=";", quotechar='"') if row}


def _sheet_rows(file_io):
    file_io.seek(0)
    sheet = load_workbook(file_io, read_only=True).active
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


# FIXTURES DE DONNÉES ==============================================================================


def test_fixtures_cree_les_centrales_et_les_enseignes(make_sale_price_category):
    """Centrale mère, cinq centrales filles et cinq enseignes depuis les fichiers csv"""
    make_sale_price_category("generique")

    fixtures()

    assert PrincipalCenterPurchase.objects.get(code="HA").name == "Héron Achats"
    assert set(ChildCenterPurchase.objects.values_list("code", flat=True)) == _csv_codes("centrales_filles.csv")
    assert {"ACF", "MAA", "GAF", "ARI", "UNI"} <= set(Signboard.objects.values_list("code", flat=True))
    assert ChildCenterPurchase.objects.get(code="GAF").base_center.code == "HA"
    assert Signboard.objects.get(code="UNI").name == "Unisson"
    assert Signboard.objects.get(code="ACF").sale_price_category.name == "generique"

    fixtures()

    assert Signboard.objects.count() == len(_csv_codes("enseignes.csv"))


def test_fixtures_journalise_les_conflits(make_sale_price_category, make_principal_center, make_child_center, make_signboard, caplog):
    """Un code déjà pris avec d'autres valeurs provoque un conflit journalisé, le chargement continue"""
    make_sale_price_category("generique")
    make_principal_center("HA", name="Autre nom")
    make_child_center("ACF", name="Autre fille")
    make_signboard("MAA", name="Autre enseigne")

    with caplog.at_level(logging.ERROR, logger="imports"):
        fixtures()

    assert PrincipalCenterPurchase.objects.get(code="HA").name == "Autre nom"
    assert ChildCenterPurchase.objects.get(code="ACF").name == "Autre fille"
    assert Signboard.objects.get(code="MAA").name == "Autre enseigne"
    assert ChildCenterPurchase.objects.count() == len(_csv_codes("centrales_filles.csv"))
    assert sum("duplicate key" in record.getMessage() for record in caplog.records) == 3


# EXPORTS EXCEL ====================================================================================


@pytest.fixture
def centrales(make_principal_center, make_child_center, make_signboard):
    mere = make_principal_center("HA", name="Héron Achats", comment="mère")
    fille = make_child_center("ACF", name="Acuitis", base_center=mere, legal_notice="Mentions", footer="Pied")
    make_signboard("ACF", name="Acuitis", child_center=fille, comment="enseigne")
    return {"mere": mere, "fille": fille}


def test_export_centrales_meres(centrales):
    file_io = io.BytesIO()

    result = output_excel_meres_list.excel_meres_list(file_io, "LISTING_DES_CENTRALES_MERES_2026_10_11_1.xlsx", PrincipalCenterPurchase)

    rows = _sheet_rows(file_io)
    assert result == {"OK": "GENERATION DU FICHIER LISTING_DES_CENTRALES_MERES_2026_10_11_1.xlsx TERMINEE AVEC SUCCES"}
    assert rows[0][0] == "LISTING DES CENTRALES MERES"
    assert rows[4][:4] == ("HA", "Héron Achats", 1, "mère")
    assert len(columns_list_meres) == 4


def test_export_centrales_filles(centrales):
    file_io = io.BytesIO()

    result = output_excel_filles_list.excel_filles_list(file_io, "LISTING_DES_CENTRALES_FILLES_2026_10_11_1.xlsx", ChildCenterPurchase)

    rows = _sheet_rows(file_io)
    assert "OK" in result
    assert rows[4][:3] == ("ACF", "Acuitis", "HA -Héron Achats")
    assert rows[4][4:6] == ("Mentions", "Pied")
    assert len(columns_list_filles) == 7


def test_export_enseignes(centrales):
    file_io = io.BytesIO()

    result = output_excel_enseignes_list.excel_enseignes_list(file_io, "LISTING_DES_ENSEIGNES_2026_10_11_1.xlsx", Signboard)

    rows = _sheet_rows(file_io)
    assert "OK" in result
    assert rows[4][:2] == ("ACF", "Acuitis")
    assert rows[4][3:5] == ("acf@heron.local", "enseigne")
    assert len(columns_list_enseignes) == 5


def test_export_regroupements(make_grouping_goods):
    make_grouping_goods("Verres", ranking=2)
    make_grouping_goods("Montures", ranking=1)
    file_io = io.BytesIO()

    result = output_excel_grouping_goods_list.excel_liste_grouping_goods(file_io, "LISTING_DES_REGROUPEMENTS_DE_FACTURATION_2026_10_11_1.xlsx")

    rows = _sheet_rows(file_io)
    assert "OK" in result
    assert output_excel_grouping_goods_list.get_clean_rows() == [("1", "Marchandises", "Montures"), ("2", "Marchandises", "Verres")]
    assert rows[4][:3] == ("1", "Marchandises", "Montures")


def test_export_axe_pro_regroupement(make_axe_pro_grouping):
    link = make_axe_pro_grouping("OPT")
    file_io = io.BytesIO()

    result = output_excel_axe_grouping_list.excel_liste_axe_grouping(file_io, "LISTING_DES_AXE_PRO_VS_REGROUPEMENTS_DE_FACTURATION_2026_10_11_1.xlsx")

    rows = _sheet_rows(file_io)
    assert "OK" in result
    assert output_excel_axe_grouping_list.get_clean_rows() == [(str(link.axe_pro), "1 - Marchandises - Montures")]
    assert rows[4][0] == str(link.axe_pro)


def test_export_axe_pro_tva(make_applicable_pro_vat):
    link = make_applicable_pro_vat("ACF", "OPT", "001")
    file_io = io.BytesIO()

    result = output_excel_axe_pro_vat_list.excel_liste_axe_vat(file_io, "LISTING_DES_CENTRALE_AXE_PRO_TVA_2026_10_11_1.xlsx")

    rows = _sheet_rows(file_io)
    assert "OK" in result
    assert output_excel_axe_pro_vat_list.get_clean_rows() == [(str(link.child_center), str(link.axe_pro), str(link.vat))]
    assert rows[4][0] == str(link.child_center)


def test_export_comptes_par_centrale(make_accounts_axe_pro_category, make_sub_category):
    """Toutes les centrales avec *, sinon seules celles de l'utilisateur (séparées par |)"""
    make_accounts_axe_pro_category(child_center="ACF", vat="001")
    make_accounts_axe_pro_category(child_center="GAF", vat="001", sub_category=make_sub_category("Montures"))
    make_accounts_axe_pro_category(child_center="MAA", vat="001")

    rows = output_excel_account_axe_list.get_clean_rows("*")
    filtered = output_excel_account_axe_list.get_clean_rows("ACF|GAF")
    file_io = io.BytesIO()
    result = output_excel_account_axe_list.excel_liste_account_axe(file_io, "LISTING_DES_COMPTES_2026_10_11_1.xlsx", "MAA")

    assert len(rows) == 3
    assert len(rows[0]) == len(output_excel_account_axe_list.columns) == 14
    assert [row[0] for row in filtered] == ["ACF", "GAF"]
    assert filtered[0][4:6] == ("", None)
    assert filtered[1][5] == "Montures"
    assert filtered[0][10:14] == ("FRA", "607000", "FRA", "707000")
    assert "OK" in result
    assert _sheet_rows(file_io)[4][0] == "MAA"
