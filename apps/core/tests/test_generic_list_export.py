"""Tests de l'export Excel générique de liste (apps.core.excel_outputs.generic_list_export)"""

import io
import logging

import openpyxl
import pytest

from apps.core.excel_outputs import generic_list_export
from apps.core.excel_outputs.excel_writer import f_entetes, f_ligne
from apps.core.excel_outputs.generic_list_export import KO_MESSAGE, excel_list_export, ok_message
from conftest import patched_attrs

COLUMNS = [
    {
        "entete": "Code",
        "f_entete": {**f_entetes, "bg_color": "#dce7f5"},
        "f_ligne": {**f_ligne, "align": "center"},
        "width": 10,
    },
    {
        "entete": "Libellé",
        "f_entete": {**f_entetes, "bg_color": "#dce7f5"},
        "f_ligne": {**f_ligne, "align": "left"},
        "width": 40,
    },
]
ROWS = [("A", "Alpha"), ("B", "Bêta"), ("C", "Gamma")]


def read_sheet(file_io):
    file_io.seek(0)
    workbook = openpyxl.load_workbook(file_io)
    sheet = workbook.worksheets[0]
    return workbook, sheet, [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_export_standard():
    file_io = io.BytesIO()

    result = excel_list_export(
        file_io, "liste.xlsx", sheet_name="LISTE", title="MA LISTE", columns=COLUMNS, rows=ROWS
    )

    assert result == ok_message("liste.xlsx")
    workbook, sheet, cells = read_sheet(file_io)
    assert workbook.sheetnames == ["LISTE"]
    assert cells[0][0] == "MA LISTE"
    assert cells[3][:2] == ["Code", "Libellé"]
    assert [row[:2] for row in cells[4:7]] == [list(row) for row in ROWS]
    assert sheet.page_setup.orientation == "portrait"


def test_export_rows_callable_et_paysage():
    file_io = io.BytesIO()
    calls = []

    def rows():
        calls.append(1)
        return iter(ROWS)

    result = excel_list_export(
        file_io,
        "liste.xlsx",
        sheet_name="L",
        title="T",
        columns=COLUMNS,
        rows=rows,
        orientation="landscape",
    )

    assert "OK" in result
    assert calls == [1]
    _, sheet, cells = read_sheet(file_io)
    assert sheet.page_setup.orientation == "landscape"
    assert cells[4][0] == "A"


def test_export_entetes_fusionnees_et_decalage():
    file_io = io.BytesIO()

    result = excel_list_export(
        file_io,
        "liste.xlsx",
        sheet_name="L",
        title="T",
        columns=COLUMNS,
        rows=ROWS,
        header_row=4,
        merged_headers=[(0, 1, "Groupe", COLUMNS[0]["f_entete"])],
        in_memory=True,
    )

    assert "OK" in result
    _, sheet, cells = read_sheet(file_io)
    assert cells[3][0] == "Groupe"
    assert cells[4][:2] == ["Code", "Libellé"]
    assert cells[5][0] == "A"
    assert "A1:B1" in [str(r) for r in sheet.merged_cells.ranges] or any(
        "A4" in str(r) for r in sheet.merged_cells.ranges
    )


@pytest.mark.parametrize("odd_first", [False, True])
def test_export_alternance_des_formats(odd_first):
    """odd_first inverse les formats pair/impair, la couleur impaire est paramétrable"""
    file_io = io.BytesIO()

    excel_list_export(
        file_io,
        "liste.xlsx",
        sheet_name="L",
        title="T",
        columns=COLUMNS,
        rows=ROWS,
        odd_color="#F2F2F2",
        odd_first=odd_first,
    )

    _, sheet, _ = read_sheet(file_io)
    first, second = sheet.cell(row=5, column=1), sheet.cell(row=6, column=1)
    greys = {c.fill.fgColor.rgb for c in (first, second)}
    assert any(rgb and rgb.endswith("F2F2F2") for rgb in greys)
    assert (first.fill.fgColor.rgb or "").endswith("F2F2F2") is odd_first


def test_export_erreur_requete_renvoie_ko(caplog):
    """Une erreur dans la fonction de lignes est journalisée et donne le retour KO"""

    def rows():
        raise RuntimeError("requête impossible")

    with caplog.at_level(logging.ERROR, logger="export_excel"):
        result = excel_list_export(
            io.BytesIO(), "liste.xlsx", sheet_name="L", title="T", columns=COLUMNS, rows=rows
        )

    assert result == KO_MESSAGE
    assert result is not KO_MESSAGE
    assert "'liste.xlsx'" in caplog.text


def test_export_erreur_ecriture_renvoie_ko(monkeypatch, caplog):
    def boom(*args, **kwargs):
        raise ValueError("boom")

    with (
        patched_attrs(monkeypatch, generic_list_export, titre_page_writer=boom),
        caplog.at_level(logging.ERROR, logger="export_excel"),
    ):
        result = excel_list_export(
            io.BytesIO(), "liste.xlsx", sheet_name="L", title="T", columns=COLUMNS, rows=ROWS
        )

    assert result == KO_MESSAGE


def test_export_hide_zero():
    """Les zéros peuvent être masqués sur la feuille"""
    file_io = io.BytesIO()

    excel_list_export(
        file_io, "liste.xlsx", sheet_name="L", title="T", columns=COLUMNS, rows=ROWS, hide_zero=True
    )

    _, sheet, _ = read_sheet(file_io)
    assert sheet.sheet_view.showZeros is False


# EXTENSIONS (phase 2) =============================================================================
from apps.core.excel_outputs.generic_list_export import (  # noqa: E402
    PAGE_NUMBERS_FOOTER,
    SheetSpec,
    excel_sheets_export,
    month_headers,
    with_headers,
    write_alternate_rows,
)


def test_with_headers_copie_sans_modifier_l_original():
    columns = with_headers(COLUMNS, {1: "Nom"})

    assert [c["entete"] for c in columns] == ["Code", "Nom"]
    assert COLUMNS[1]["entete"] == "Libellé"
    assert columns[0] is not COLUMNS[0]


def test_month_headers():
    assert [c["entete"] for c in month_headers(COLUMNS, (0, 1), ["M-1", "M"])] == ["M-1", "M"]


def test_columns_callable_et_writer_sans_entetes():
    """Colonnes calculées, écriture spécifique des lignes, pas d'entêtes automatiques, pied de page"""
    file_io = io.BytesIO()
    seen = {}

    def writer(excel, sheet, columns, rows, f_lignes, f_lignes_odd):
        seen["columns"] = [c["entete"] for c in columns]
        count = write_alternate_rows(excel, sheet, 4, rows, f_lignes, f_lignes_odd)
        excel.write_row(sheet, 4 + count, 0, "FIN", f_lignes[0])
        return [4 + count]

    result = excel_list_export(
        file_io,
        "liste.xlsx",
        sheet_name="L",
        title="T",
        columns=lambda: with_headers(COLUMNS, {0: "Réf"}),
        rows=ROWS,
        write_headers=False,
        writer=writer,
        footer=PAGE_NUMBERS_FOOTER,
    )

    assert "OK" in result
    assert seen["columns"] == ["Réf", "Libellé"]
    _, sheet, cells = read_sheet(file_io)
    assert cells[3][0] is None  # pas d'entêtes de colonnes
    assert cells[4][0] == "A"
    assert cells[7][0] == "FIN"
    assert sheet.HeaderFooter.oddFooter.right.text == "&P/&N" or "&P" in str(
        sheet.oddFooter.right.text
    )
    assert [b.id for b in sheet.row_breaks.brk] == [7]


def test_after_rows_recoit_la_plage_ecrite():
    calls = []

    excel_list_export(
        io.BytesIO(),
        "liste.xlsx",
        sheet_name="L",
        title="T",
        columns=COLUMNS,
        rows=ROWS,
        after_rows=lambda excel, sheet, first_row, nb_rows: calls.append(
            (sheet, first_row, nb_rows)
        ),
    )

    assert calls == [(1, 4, 3)]


def test_excel_sheets_export_plusieurs_feuilles():
    file_io = io.BytesIO()
    specs = [
        SheetSpec(name="UNE", title="Première", columns=COLUMNS, rows=ROWS[:1]),
        SheetSpec(
            name="DEUX",
            title="Seconde",
            columns=COLUMNS,
            rows=ROWS,
            extra_formatting={"not_display_zero": True},
            hide_zero=True,
        ),
    ]

    result = excel_sheets_export(file_io, "classeur.xlsx", specs, in_memory=True)

    assert result == ok_message("classeur.xlsx")
    file_io.seek(0)
    workbook = openpyxl.load_workbook(file_io)
    assert workbook.sheetnames == ["UNE", "DEUX"]
    assert workbook["UNE"]["A1"].value == "Première"
    assert workbook["DEUX"]["A5"].value == "A"
    assert workbook["DEUX"]["A7"].value == "C"
    assert workbook["DEUX"].sheet_view.showZeros is False
