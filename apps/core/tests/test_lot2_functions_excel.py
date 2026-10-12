# pylint: disable=E0401,W0621
"""
FR : Tests des utilitaires Excel (functions_excel, functions_excel_optimized, excel_outputs/excel_writer)
EN : Tests for Excel utilities

created at: 2026-10-11
created by: Paulo ALVES
"""
import csv
import datetime
import glob
import zipfile
from pathlib import Path

import pytest
import xlsxwriter
from openpyxl import load_workbook

from apps.core.excel_outputs import excel_writer
from apps.core.excel_outputs.excel_writer import (
    columns_headers_writer,
    f_entetes,
    f_ligne,
    f_titre,
    output_day_writer,
    sheet_formatting,
    titre_page_writer,
)
from apps.core.functions import functions_excel
from apps.core.functions.functions_excel import (
    XLSX_PAPER_SIZE,
    ExcelToCsv,
    ExcelToCsvError,
    GenericExcel,
    date_debut_nom_fichier,
    excel_excel_to_csv,
    excel_xhtml_to_csv,
    get_paper_size,
)
from apps.core.functions.functions_excel_optimized import GenericExcelOptimized
from conftest import patched_attrs

FIXTURES = Path(__file__).resolve().parent / "test_fixtures"
XHTML = (
    "<html><body><table>"
    "<tr><td><span>Tiers</span></td><td><span>Montant=;</span></td><td></td></tr>"
    "<tr><td><span>FOUR001</span></td><td><span>12,5\"</span></td><td><span>x\n\r</span></td></tr>"
    "</table></body></html>"
)


def _read_csv(path):
    with Path(path).open("r", encoding="utf8", newline="") as file:
        return list(csv.reader(file, delimiter=";", quotechar='"'))


def _xlsx(path, rows, sheets=("Feuil1",)):
    """Classeur avec les mêmes lignes sur chaque feuille"""
    workbook = xlsxwriter.Workbook(str(path))

    for sheet in sheets:
        worksheet = workbook.add_worksheet(sheet)

        for row_index, row in enumerate(rows):
            worksheet.write_row(row_index, 0, row)

    workbook.close()
    return path


@pytest.fixture
def no_log(monkeypatch):
    """Capture les écritures de log de functions_excel au lieu d'écrire dans files/"""
    lines = []

    with patched_attrs(monkeypatch, functions_excel, write_log=lambda fichier, ligne: lines.append((fichier, ligne))):
        yield lines


# CONVERSIONS ======================================================================================


def test_get_paper_size():
    assert get_paper_size("A4") == 9
    assert get_paper_size("Inconnu") == 0
    assert all(isinstance(code, int) for code, _ in XLSX_PAPER_SIZE.values())


def test_excel_excel_to_csv_une_feuille(tmp_path):
    source = _xlsx(tmp_path / "classeur.xlsx", [["Tiers", "Montant"], ["FOUR001", 12.5]])
    target = tmp_path / "classeur.csv"

    assert excel_excel_to_csv(source, target) == [target]
    assert _read_csv(target) == [["Tiers", "Montant"], ["FOUR001", "12.5"]]
    assert excel_excel_to_csv(source, target, header=False) == [target]
    assert _read_csv(target) == [["FOUR001", "12.5"]]


def test_excel_excel_to_csv_plusieurs_feuilles(tmp_path):
    """Un csv par feuille, nommé d'après la feuille (sans espace, en majuscules)"""
    source = _xlsx(tmp_path / "classeur.xlsx", [["a", "b"], [1, 2]], sheets=("Feuil 1", "Deux"))

    result = excel_excel_to_csv(source, str(tmp_path / "ignore.csv"), sheets=["Feuil 1", "Deux"])

    assert result == [str(tmp_path / "FEUIL1.csv"), str(tmp_path / "DEUX.csv")]
    assert _read_csv(result[1]) == [["a", "b"], ["1", "2"]]


def test_excel_excel_to_csv_ancien_format_xls(tmp_path):
    """Un classeur .xls (BIFF) est lu par xlrd"""
    target = tmp_path / "ancien.csv"

    assert excel_excel_to_csv(FIXTURES / "test_pandas.xls", target) == [target]
    assert _read_csv(target) == [["test", "pandas", "1", "=67"]]


def test_excel_excel_to_csv_fichier_non_reconnu_tente_openpyxl(tmp_path):
    """Format indéterminable : second essai avec openpyxl, qui refuse un fichier non zip"""
    source = tmp_path / "faux.xlsx"
    source.write_text("pas un classeur", encoding="utf8")

    with pytest.raises(zipfile.BadZipFile):
        excel_excel_to_csv(source, tmp_path / "faux.csv")


def test_excel_excel_to_csv_erreur_module_journalisee(tmp_path, no_log, monkeypatch):
    """L'erreur propre au module est journalisée et aucun fichier n'est renvoyé"""

    def _boom(*_args, **_kwargs):
        raise ExcelToCsvError("conversion impossible")

    with patched_attrs(monkeypatch, functions_excel.pd, ExcelFile=_boom):
        assert excel_excel_to_csv(tmp_path / "x.xlsx", tmp_path / "x.csv") == []

    assert "Erreur excel_excel_to_csv : conversion impossible" in no_log[0][1]


def test_excel_xhtml_to_csv(tmp_path):
    """Les cellules (span) sont nettoyées des =, guillemets, ; et retours à la ligne"""
    source = tmp_path / "export.xls"
    source.write_text(XHTML, encoding="utf8")
    target = tmp_path / "export.csv"

    assert excel_xhtml_to_csv(source, target) == [str(target)]
    assert _read_csv(target) == [["Tiers", "Montant", ""], ["FOUR001", "12,5", "x"]]

    assert excel_xhtml_to_csv(source, target, header=None) == [str(target)]
    assert _read_csv(target) == [["FOUR001", "12,5", "x"]]


def test_excel_xhtml_to_csv_erreur_module_journalisee(tmp_path, no_log, monkeypatch):
    def _boom(_path):
        raise ExcelToCsvError("encodage")

    with patched_attrs(monkeypatch, functions_excel, encoding_detect=_boom):
        assert excel_xhtml_to_csv(tmp_path / "export.xls", tmp_path / "export.csv") == []

    assert "Erreur excel_excel_to_csv : encodage" in no_log[0][1]


def test_date_debut_nom_fichier():
    assert date_debut_nom_fichier() == datetime.date.today().isoformat().replace("-", "") + "_"


# CLASSE EXCEL TO CSV ==============================================================================


def test_excel_to_csv_sans_repertoire():
    assert ExcelToCsv().make_csv() == []
    assert ExcelToCsv(log="journal.log").log_file == "journal.log"


def test_excel_is_xhtml(tmp_path):
    xhtml = tmp_path / "a.xls"
    xhtml.write_text(XHTML, encoding="utf8")
    _xlsx(tmp_path / "b.xlsx", [["a"]])

    assert ExcelToCsv.excel_is_xhtml(xhtml) is True
    assert ExcelToCsv.excel_is_xhtml(tmp_path / "b.xlsx") is False


def test_make_csv_convertit_tous_les_classeurs(tmp_path):
    """Classeur xlsx, export xhtml nommé .xls, fichier texte ignoré ; date en préfixe des csv"""
    _xlsx(tmp_path / "classeur.xlsx", [["a", "b"], [1, 2]])
    (tmp_path / "export.xls").write_text(XHTML, encoding="utf8")
    (tmp_path / "notes.txt").write_text("ignoré", encoding="utf8")
    prefix = date_debut_nom_fichier()

    treated = ExcelToCsv(rep=str(tmp_path)).make_csv()

    assert sorted(Path(file).name for file in treated) == [f"{prefix}classeur.csv", f"{prefix}export.csv"]
    assert _read_csv(tmp_path / f"{prefix}export.csv") == [["Tiers", "Montant", ""], ["FOUR001", "12,5", "x"]]
    assert (tmp_path / "classeur.xlsx").exists()


def test_make_csv_un_fichier_sans_date_avec_suppression(tmp_path):
    """Un seul fichier ciblé, pas de préfixe de date, source supprimée après conversion"""
    _xlsx(tmp_path / "classeur.xlsx", [["a"], [1]])
    _xlsx(tmp_path / "autre.xlsx", [["b"], [2]])
    (tmp_path / "classeur.csv").write_text("ancien", encoding="utf8")

    treated = ExcelToCsv(rep=str(tmp_path), file_name="classeur.xlsx", first_date=None, deletion=True, header=False).make_csv()

    assert treated == [str(tmp_path / "classeur.csv")]
    assert _read_csv(treated[0]) == [["1"]]
    assert not (tmp_path / "classeur.xlsx").exists()
    assert (tmp_path / "autre.xlsx").exists()


def test_make_csv_deplace_la_source_en_sauvegarde(tmp_path):
    source_dir = tmp_path / "source"
    backup_dir = tmp_path / "backup"
    source_dir.mkdir()
    backup_dir.mkdir()
    _xlsx(source_dir / "classeur.xlsx", [["a"], [1]])

    treated = ExcelToCsv(rep=str(source_dir), rep_sauv=str(backup_dir), first_date=None).make_csv()

    assert treated == [str(source_dir / "classeur.csv")]
    assert (backup_dir / "classeur.xlsx").exists()
    assert not (source_dir / "classeur.xlsx").exists()


def test_make_csv_conversion_sans_resultat_journalisee(tmp_path, monkeypatch):
    """Si la conversion ne renvoie rien, l'échec est journalisé dans le fichier de log donné"""
    _xlsx(tmp_path / "classeur.xlsx", [["a"], [1]])
    log = tmp_path / "journal.log"

    with patched_attrs(monkeypatch, functions_excel, excel_excel_to_csv=lambda *_a, **_k: []):
        treated = ExcelToCsv(rep=str(tmp_path), log=str(log)).make_csv()

    assert treated == []
    assert "Erreur de Conversion .xls en .csv : classeur.xlsx" in log.read_text(encoding="utf8")


def test_make_csv_erreur_module_journalisee(tmp_path, monkeypatch):
    def _boom(_pattern):
        raise ExcelToCsvError("répertoire illisible")

    log = tmp_path / "journal.log"

    with patched_attrs(monkeypatch, glob, glob=_boom):
        assert ExcelToCsv(rep=str(tmp_path), log=str(log)).make_csv() == []

    assert "Erreur de conversion fichiers .xls en .csv : répertoire illisible" in log.read_text(encoding="utf8")


# GENERIC EXCEL ====================================================================================


def _sheet(path, name="Feuil1"):
    return load_workbook(path, read_only=False)[name]


def test_generic_excel_ecritures(tmp_path):
    """Cellules simples, lignes, fusions, entêtes à deux étages, image et options de feuille"""
    path = tmp_path / "classeur.xlsx"
    bold = {"bold": True}
    excel = GenericExcel([str(path), ["Feuil1", "Feuil2"]])

    excel.write_row(1, 0, 0, "titre")
    excel.write_row(1, 0, 1, "gras", bold)
    excel.write_rows(1, 1, 0, ["a", "b", "c"])
    excel.write_rows(1, 2, 0, ["d", "e"], bold)
    excel.write_rows(1, 3, 0, ["f", "g"], [bold, {"italic": True}])
    excel.write_rows(1, 4, 0, ["h", "i", "j"], [bold])
    excel.write_merge_h(1, 5, 0, 2, "fusion h")
    excel.write_merge_h(1, 6, 0, 2, "fusion h gras", bold)
    excel.write_merge_v(1, 7, 9, 0, "fusion v")
    excel.write_merge_v(1, 7, 9, 1, "fusion v gras", bold)
    excel.write_merge_v(1, 10, 10, 0, "une cellule")
    excel.write_merge_v(1, 10, 10, 1, "une cellule gras", bold)
    excel.write_merge(1, "D1:E2", "plage")
    excel.write_merge(1, "D3:E4", "plage gras", bold)
    excel.write_headers_h(2, 0, 0, ["Entête", "t1", "t2", "t3"])
    excel.write_headers_h(2, 3, 0, ["Entête", "t1", "t2"], [bold, bold, bold])
    excel.write_headers_v(2, 6, 0, ["E", "l1", "l2", "l3"])
    excel.write_headers_v(2, 6, 3, ["E", "l1", "l2"], [bold, bold, bold])
    excel.sheet_zoom(1, 80)
    excel.sheet_hide_zero(1)
    excel.sheet_tab_color(1, "red")
    excel.conditional_value(1, "A1:C4", 0)
    excel.conditional_value(1, "A1:C4", 0, style=bold)
    assert excel.excel_sheet(2).name == "Feuil2"
    excel.excel_close()

    sheet = _sheet(path)
    second = _sheet(path, "Feuil2")
    assert [cell.value for cell in sheet[1]][:2] == ["titre", "gras"]
    assert [cell.value for cell in sheet[2]][:3] == ["a", "b", "c"]
    assert sheet["A3"].font.bold is True
    assert sheet["B4"].font.italic is True
    assert sheet["C5"].font.bold is True
    assert sheet["A6"].value == "fusion h"
    assert sheet["A8"].value == "fusion v"
    assert sheet["A11"].value == "une cellule"
    assert sheet["D1"].value == "plage"
    assert {str(rng) for rng in sheet.merged_cells.ranges} >= {"A6:C6", "A8:A10", "D1:E2"}
    assert second["A1"].value == "Entête"
    assert [second["A2"].value, second["B2"].value, second["C2"].value] == ["t1", "t2", "t3"]
    assert second["A7"].value == "E"
    assert [second["B7"].value, second["B8"].value, second["B9"].value] == ["l1", "l2", "l3"]
    assert sheet.sheet_properties.tabColor.rgb.endswith("FF0000")


def test_generic_excel_image_et_memoire_constante(tmp_path):
    from PIL import Image  # pylint: disable=C0415

    image = tmp_path / "logo.png"
    Image.new("RGB", (4, 4), "blue").save(image)
    path = tmp_path / "classeur.xlsx"
    excel = GenericExcel([str(path), ["Feuil1"]], in_memory=True)

    excel.write_image(1, 0, 0, str(image))
    excel.write_image(1, 5, 0, str(image), {"x_scale": 0.5})
    excel.excel_close()

    assert len(_sheet(path)._images) == 2  # pylint: disable=W0212


# GENERIC EXCEL OPTIMISÉ ===========================================================================


def test_generic_excel_optimized_cache_des_formats(tmp_path):
    """Un format par style distinct, réutilisé entre les écritures"""
    path = tmp_path / "classeur.xlsx"
    bold = {"bold": True}
    excel = GenericExcelOptimized([str(path), ["Feuil1"]])

    excel.write_row(1, 0, 0, "sans style")
    excel.write_row(1, 0, 1, "gras", bold)
    excel.write_rows(1, 1, 0, ["a", "b"])
    excel.write_rows(1, 2, 0, ["c", "d"], bold)
    excel.write_rows(1, 3, 0, ["e", "f"], [bold, {"italic": True}])
    excel.write_rows(1, 4, 0, ["g", "h", "i"], [bold])
    excel.write_rows(1, 5, 0, ["j", "k"], [])
    formats = excel.prepare_row_formats([bold, None])
    excel.write_rows_with_prepared_formats(1, 6, 0, ("l", "m", "n"), formats)
    stats = excel.get_cache_stats()
    assert excel._get_format_key(None) == "__none__"  # pylint: disable=W0212
    assert excel._get_cached_format(None) is None  # pylint: disable=W0212
    excel.excel_close()

    sheet = _sheet(path)
    assert stats == {"formats_cached": 2}
    assert formats[1] is None
    assert [cell.value for cell in sheet[5]][:3] == ["g", "h", "i"]
    assert sheet["A5"].font.bold is True
    assert sheet["C5"].font.bold is True
    assert sheet["A6"].font.bold is not True
    assert sheet["A7"].font.bold is True
    assert sheet["B7"].font.bold is not True
    assert sheet["C7"].value == "n"


# EXCEL WRITER =====================================================================================

COLUMNS = [
    {"entete": "Colonne 1", "f_entete": f_entetes, "f_ligne": f_ligne, "width": 10},
    {"entete": "Colonne 2", "f_entete": f_entetes, "f_ligne": f_ligne, "width": 20},
    {"entete": "Colonne 3", "f_entete": f_entetes, "f_ligne": f_ligne, "width": 30},
]


def test_excel_writer_titre_date_et_entetes(tmp_path):
    """Titre sur toute la largeur (fusionné ou non), date d'édition, entêtes de colonnes"""
    path = tmp_path / "classeur.xlsx"
    excel = GenericExcelOptimized([str(path), ["Feuil1"]])

    titre_page_writer(excel, 1, 0, 0, COLUMNS, "Titre non fusionné")
    titre_page_writer(excel, 1, 1, 0, COLUMNS, "Titre fusionné", merge=True)
    output_day_writer(excel, 1, 2, 0)
    output_day_writer(excel, 1, 3, 0, {"italic": True})
    columns_headers_writer(excel, 1, 4, 0, COLUMNS)
    excel.excel_close()

    sheet = _sheet(path)
    assert [cell.value for cell in sheet[1]][:3] == ["Titre non fusionné", None, None]
    assert sheet["A2"].value == "Titre fusionné"
    assert "A2:D2" in {str(rng) for rng in sheet.merged_cells.ranges}
    assert f_titre["align"] == "center"
    assert sheet["A3"].value.startswith("Edité, le ")
    assert sheet["A4"].font.italic is True
    assert [cell.value for cell in sheet[5]][:3] == ["Colonne 1", "Colonne 2", "Colonne 3"]


@pytest.mark.parametrize(
    "dict_args",
    [
        None,
        {"sens": "landscape", "centrer": False, "repeat_row": (0, 2), "fit_page": (1, 0), "not_display_zero": True, "paper": "A3", "marge": 0.5},
    ],
)
def test_excel_writer_mise_en_page(tmp_path, dict_args):
    """Largeurs de colonnes, papier, marges, orientation et options d'impression"""
    path = tmp_path / "classeur.xlsx"
    excel = GenericExcelOptimized([str(path), ["Feuil1"]])

    sheet_formatting(excel, 1, COLUMNS, dict_args)
    worksheet = excel.excel_sheet(1)
    excel.excel_close()

    assert worksheet.paper_size == (8 if dict_args else 9)
    assert worksheet.orientation == (0 if dict_args else 1)
    assert worksheet.hcenter == (0 if dict_args else 1)
    assert worksheet.repeat_row_range == ("Feuil1!$1:$3" if dict_args else "")
    assert worksheet.col_info[2][0] == 30


def test_excel_writer_formats_par_defaut():
    assert excel_writer.f_edit["bold"] is True
    assert f_entetes["bg_color"] == "#dce7f5"
    assert f_ligne["valign"] == "top"
