# pylint: disable=E0401,W0621,W0212
"""FR : Tests de la configuration, des modèles, des formulaires, des fixtures et des exports Excel


     de l'app parameters
EN : Tests for parameters app configuration, models, forms, fixtures and Excel exports

created at: 2026-10-11
created by: Paulo ALVES
"""

import importlib
import io
from decimal import Decimal

import pytest
from django.apps import apps as django_apps
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.urls import reverse
from openpyxl import load_workbook

from apps.parameters import urls as parameters_urls
from apps.parameters.apps import ParametersConfig
from apps.parameters.data_fixtures.insert_fixtures import fixtures
from apps.parameters.excel_outputs import (
    parameters_excel_axe_articles_defaut_list,
    parameters_excel_categories_list,
    parameters_excel_control_rebilling_list,
    parameters_excel_functions,
    parameters_excel_natures,
    parameters_excel_numberings,
)
from apps.parameters.forms import (
    CounterForm,
    DeleteControlRebillingForm,
    DeleteSubCategoryForm,
    EmailForm,
    ExchangeRateForm,
    NumberInput,
)
from apps.parameters.models import (
    ActionInProgress,
    ActionPermission,
    AddressCode,
    Category,
    Counter,
    CounterNums,
    DefaultAxeArticle,
    Email,
    ExchangeRate,
    IconOriginChoice,
    InvoiceFunctions,
    Nature,
    Parameters,
    Periodicity,
    SalePriceCategory,
    SendFiles,
    SendFilesMail,
    SubCategory,
    SubFamilly,
)

pytestmark = pytest.mark.django_db


def _sheet_rows(file_io):
    file_io.seek(0)
    return [
        tuple(row)
        for row in load_workbook(file_io, read_only=True).active.iter_rows(values_only=True)
    ]


# CONFIGURATION ====================================================================================


def test_app_config_et_routes():
    config = django_apps.get_app_config("parameters")
    names = [pattern.name for pattern in parameters_urls.urlpatterns]

    assert isinstance(config, ParametersConfig)
    assert config.name == "apps.parameters"
    assert len(names) == 36
    assert len(set(names)) == 36
    assert len(list(config.get_models())) == 21


@pytest.mark.parametrize(
    "module_name",
    [
        "apps.parameters.admin",
        "apps.parameters.parameters",
        "apps.parameters.forms.forms_django.const_forms",
        "apps.parameters.bin",
    ],
)
def test_packages_importables(module_name):
    assert importlib.import_module(module_name) is not None


# MODÈLES ==========================================================================================


def test_affichages_simples(
    make_unit, make_nature, make_sale_price_category, make_invoice_function, make_user
):
    parameter = Parameters.objects.create(name="taux", value="2")
    send_file = SendFiles.objects.create(name="Export", file="export.xlsx")
    mail = SendFilesMail.objects.create(file=send_file, user=make_user(), email="a@b.fr")

    assert str(parameter) == "taux"
    assert str(make_unit(1, "U")) == "U"
    assert str(SubFamilly.objects.create(name="Montures")) == "Montures"
    assert str(Periodicity.objects.create(name="Mensuel")) == "Mensuel"
    assert str(make_sale_price_category("generique")) == "generique"
    assert ActionPermission.objects.create(name="valider").name == "valider"
    assert str(make_invoice_function("ABONNEMENT")) == "ABONNEMENT"
    assert str(send_file) == "Export"
    assert str(mail) == f"Export - {mail.user} - a@b.fr"
    assert str(make_nature("Monsieur", to_display="M.")) == "M."
    assert str(make_nature("Sans libellé", to_display=None)) == "Sans libellé"
    assert Nature.get_absolute_url() == reverse("parameters:natures_list")
    assert AddressCode._meta.fields[0].name == "id"
    assert IconOriginChoice.objects.create(origin=1, origin_name="Fichier", icon="file").origin == 1
    assert ActionInProgress.objects.create(action="x").in_progress is False


def test_fonction_de_facturation(make_invoice_function):
    """Le champ reverse est calculé depuis redirect quand c'est une route connue"""
    known = make_invoice_function("ABONNEMENT", redirect="parameters:functions_list")
    unknown = make_invoice_function("FORMATION", redirect="parameters:absente")
    plain = make_invoice_function("PERSONNEL", redirect="")

    assert known.reverse == reverse("parameters:functions_list")
    assert unknown.reverse == ""
    assert plain.reverse == ""
    assert (
        InvoiceFunctions.get_success_url()
        == InvoiceFunctions.get_absolute_url()
        == reverse("parameters:functions_list")
    )


def test_compteur_controles_et_affichage(make_counter, make_invoice_function):
    """Lpad négatif refusé, fonction déjà prise par un autre compteur refusée à la création"""
    function = make_invoice_function("ABONNEMENT")
    taken = make_counter("abonnements", function=function)
    counter = Counter(name="x", lpad_num=-1)

    assert str(taken) == "abonnements"
    assert Counter.get_absolute_url() == reverse("parameters:numberings_list")
    assert Counter.bool_function(None) is False
    assert Counter.bool_function(function) is True
    assert Counter.bool_function(make_invoice_function("FORMATION")) is False

    with pytest.raises(ValidationError, match="lpad_num"):
        counter.clean()

    duplicate = Counter(name="y", lpad_num=None, function=function)

    with pytest.raises(ValidationError, match="function"):
        duplicate.clean()

    taken.clean()
    Counter(name="libre", lpad_num=0).clean()
    assert CounterNums.objects.get(counter=taken).num == 1

    with pytest.raises(IntegrityError), transaction.atomic():
        Counter.objects.create(name="negatif", lpad_num=-2)


def test_categories_et_rubriques(make_big_category, make_sub_category):
    """Slug calculé à la sauvegarde, affichages, rubrique sans catégorie"""
    category = make_big_category("Marchandises Neuves")
    sub = make_sub_category("Montures", big_category=category)

    assert category.slug_name == "marchandises-neuves"
    assert str(category) == f"{category.ranking:02d} - Marchandises Neuves"
    assert str(sub) == f"Marchandises Neuves - {sub.ranking} - MONTURES - Montures"
    assert str(SubCategory(code="X", name="Orpheline", ranking=1)) == ""
    assert (
        Category.get_absolute_url()
        == SubCategory.get_absolute_url()
        == reverse("parameters:categories_list")
    )

    category.name = "Renommée"
    category.save()
    assert category.slug_name == "marchandises-neuves"

    with pytest.raises(IntegrityError), transaction.atomic():
        SubCategory.objects.create(
            code="Y", name="Doublon de rang", ranking=sub.ranking, big_category=category
        )


def test_axes_par_defaut_email_taux_et_controle(
    make_default_axes, make_email, make_exchange_rate, make_control
):
    axes = make_default_axes()
    email = make_email(1, description="Documents")
    rate = make_exchange_rate("CHF", rate="1.05")
    control = make_control("Essilor")

    assert DefaultAxeArticle.get_absolute_url() == reverse(
        "parameters:axes_articles_defaut", kwargs={"slug_name": "axes_articles"}
    )
    assert axes.slug_name == "axes_articles"
    assert Email.get_absolute_url() == reverse("parameters:emails_list")
    assert email.get_name_display() == "Documents"
    assert (rate.curency_base, rate.rate) == ("EUR", Decimal("1.05"))
    assert str(control) == "Essilor"
    assert control.get_absolute_url() == reverse("parameters:control_rebilling_list")

    negative = ExchangeRate(rate=Decimal(-1))

    with pytest.raises(ValidationError, match="positif"):
        negative.clean()

    ExchangeRate(rate=Decimal(1)).clean()

    with pytest.raises(IntegrityError), transaction.atomic():
        ExchangeRate.objects.create(
            currency_change=rate.currency_change, rate_month=rate.rate_month
        )


# FORMULAIRES ======================================================================================


def test_formulaires_specifiques(make_invoice_function, make_country_currency):
    make_country_currency("CHF", "Franc suisse")
    counter = CounterForm(
        {
            "name": "factures",
            "lpad_num": "",
            "prefix": "",
            "suffix": "",
            "function": "",
            "description": "",
            "separateur": "",
        }
    )
    rate = ExchangeRateForm({"currency_change": "CHF", "rate": "1.05"})
    email = EmailForm(
        {"description": "Factures", "subject": "Vos factures", "email_body": "<p>Bonjour</p>"}
    )

    assert NumberInput().input_type == "number"
    assert counter.is_valid(), counter.errors
    assert counter.save().lpad_num is None
    assert rate.is_valid(), rate.errors
    assert rate.fields["curency_base"].required is False
    assert email.fields["email_body"].widget.attrs["is"] == "textarea-autogrow"
    assert email.is_valid(), email.errors
    assert ExchangeRateForm({"currency_change": "XXX", "rate": "1"}).is_valid() is False


def test_formulaires_de_suppression(make_sub_category, make_control):
    sub = make_sub_category("Montures")
    control = make_control()

    assert DeleteSubCategoryForm({"id": sub.pk}).is_valid()
    assert DeleteSubCategoryForm({"id": sub.pk + 1000}).is_valid() is False
    assert DeleteControlRebillingForm({"id": control.pk}).is_valid()
    assert DeleteControlRebillingForm({"id": ""}).is_valid() is False


# FIXTURES =========================================================================================


def test_fixtures_categories_de_prix(make_sale_price_category):
    """La catégorie de prix générique est créée, un second passage ou un conflit ne change rien"""
    fixtures()
    assert SalePriceCategory.objects.get(name="generique").coefficient == 1

    fixtures()
    assert SalePriceCategory.objects.count() == 1

    SalePriceCategory.objects.all().delete()
    make_sale_price_category("generique").__class__.objects.filter(name="generique").update(
        coefficient=2
    )

    fixtures()

    assert SalePriceCategory.objects.count() == 1


# EXPORTS EXCEL ====================================================================================


def test_export_fonctions(make_invoice_function):
    make_invoice_function("ABONNEMENT", absolute_path_linux="/home", description="Abonnements")
    file_io = io.BytesIO()

    result = parameters_excel_functions.excel_liste_functions(
        file_io, "LISTING_DES_FONCTIONS_2026_10_11_1.xlsx"
    )

    assert "OK" in result
    assert list(parameters_excel_functions.get_row()) == [
        ("ABONNEMENT", "def abonnement(): ...", "", "/home", "Abonnements")
    ]
    assert _sheet_rows(file_io)[4][0] == "ABONNEMENT"


def test_export_natures(make_nature):
    make_nature("Monsieur", to_display="M.", for_contact=True)
    file_io = io.BytesIO()

    result = parameters_excel_natures.excel_liste_natures(
        file_io, "LISTING_DES_NATURES_GENRE_2026_10_11_1.xlsx"
    )

    assert "OK" in result
    assert list(parameters_excel_natures.get_row()) == [("Monsieur", "M.", "X", "", "")]
    assert _sheet_rows(file_io)[4][:3] == ("Monsieur", "M.", "X")


def test_export_numerotations(make_counter):
    make_counter("factures", prefix="FAC", lpad_num=6, description="Factures de vente")
    file_io = io.BytesIO()

    result = parameters_excel_numberings.excel_liste_numberings(
        file_io, "LISTING_DES_NUMEROTATIONS_2026_10_11_1.xlsx"
    )

    assert "OK" in result
    assert list(parameters_excel_numberings.get_row()) == [
        ("factures", None, "FAC", None, 6, "", "Factures de vente")
    ]
    assert _sheet_rows(file_io)[4][0] == "factures"


def test_export_controles(make_control):
    make_control("Essilor", comment="à vérifier")
    file_io = io.BytesIO()

    result = parameters_excel_control_rebilling_list.excel_liste_control_rebilling(
        file_io, "LISTING_DES_CONTROLES_DE_REFACTURATION_2026_10_11_1.xlsx"
    )

    assert "OK" in result
    assert list(parameters_excel_control_rebilling_list.get_row()) == [("Essilor", "à vérifier")]
    assert _sheet_rows(file_io)[5][:2] == ("Essilor", "à vérifier")


def test_export_categories(make_big_category, make_sub_category):
    category = make_big_category("Marchandises")
    make_sub_category("Montures", big_category=category)
    make_big_category("Services")
    file_io = io.BytesIO()

    result = parameters_excel_categories_list.excel_liste_categories(
        file_io, "LISTING_DES_CATEGORIES_2026_10_11_1.xlsx"
    )

    rows = list(parameters_excel_categories_list.get_row())
    assert "OK" in result
    assert len(rows) == 2
    assert rows[0][1:3] == ("MARCHANDISES", "Marchandises")
    assert rows[0][4:6] == ("MONTURES", "Montures")
    assert rows[1][3:6] == ("", "", "")
    assert _sheet_rows(file_io)[3][0] == "Grande Catégorie"


def test_export_axes_par_defaut(make_default_axes, make_section, make_sub_category):
    make_default_axes(axe_bu=make_section("BU", "CAHA"), sub_category=make_sub_category("Montures"))
    file_io = io.BytesIO()

    result = parameters_excel_axe_articles_defaut_list.excel_list_axe_article_defaut(
        file_io, "LISTING_DES_AXES_ARTICLES_PAR_DEFAUT_2026_10_11_1.xlsx"
    )

    rows = list(parameters_excel_axe_articles_defaut_list.get_row())
    assert "OK" in result
    assert len(rows) == 1
    assert "Montures" in rows[0][1]
    assert "CAHA" in rows[0][2]
    assert "NAF" in rows[0][3]
    assert _sheet_rows(file_io)[4][2] == rows[0][2]
