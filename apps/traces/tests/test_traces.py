# pylint: disable=E0401,W0621
"""FR : Tests de l'app traces (écran paginé des traces EDI et export Excel)


EN : Tests for traces app

created at: 2026-10-12
created by: Paulo ALVES
"""

import importlib
import io

import pytest
from django.apps import apps as django_apps
from django.urls import reverse
from openpyxl import load_workbook

from apps.core.functions.functions_http_response import CONTENT_TYPE_EXCEL
from apps.data_flux.models import Error, Line
from apps.traces import views as traces_views
from apps.traces.apps import TracesConfig
from apps.traces.excel_outputs.edi_traces_exports import excel_edi_traces, get_rows

pytestmark = pytest.mark.django_db


@pytest.fixture
def traces(make_trace):
    """Une trace en erreur (ligne + erreur détaillée), une trace saine sans commentaire"""
    failed = make_trace(
        trace_name="Import BBGR",
        file_name="bbgr.csv",
        errors=True,
        comment="<p>2 lignes <b>rejetées</b></p>",
    )
    line = Line.objects.create(
        trace=failed, insertion_type="Errors", num_line=12, designation="Facture 123"
    )
    Error.objects.create(
        line=line,
        attr_name="invoice_date",
        file_column="Date",
        message="Date invalide",
        data_expected="AAAA-MM-JJ",
        data_received="31/13/2026",
    )
    clean = make_trace(trace_name="Import Sage", file_name="sage.csv", comment=None)
    return {"failed": failed, "clean": clean}


def test_app_config_et_routes():
    config = django_apps.get_app_config("traces")

    assert isinstance(config, TracesConfig)
    assert config.name == "apps.traces"
    assert list(config.get_models()) == []
    assert importlib.import_module("apps.traces.admin") is not None


def test_ecran_des_traces(logged_client, traces, make_trace):
    """Pagination de 50 : page absente, non numérique ou trop grande retombe sur une page valide"""
    for i in range(55):
        make_trace(trace_name=f"trace {i}")
    url = reverse("traces:edi_traces")

    first = logged_client.get(url)
    text = logged_client.get(url, {"page": "abc"})
    last = logged_client.get(url, {"page": 99})

    assert first.status_code == 200
    assert first.context["num_items"] == 57
    assert first.context["num_pages"] == 2
    assert (first.context["start_index"], first.context["end_index"]) == (0, 50)
    assert text.context["traces"].number == 1
    assert last.context["traces"].number == 2
    assert (last.context["start_index"], last.context["end_index"]) == (50, 57)
    assert "cadetblue" in str(first.context["pagination"])
    assert "trace 54" in first.content.decode()
    assert "Import BBGR" in last.content.decode()


def test_ecran_sans_trace(logged_client):
    response = logged_client.get(reverse("traces:edi_traces"))

    assert response.status_code == 200
    assert (
        response.context["num_items"],
        response.context["start_index"],
        response.context["end_index"],
    ) == (0, 0, 0)


def test_lignes_export(traces):
    rows = get_rows(0, 50)

    assert len(rows) == 2
    clean, failed = rows
    assert clean[2:6] == ("sage.csv", "Import Sage", "", "")
    assert clean[6] == ""
    assert failed[0] == str(traces["failed"].uuid_identification)
    assert failed[4:6] == ("X", "2 lignes **rejetées**")
    assert "ligne n° : 12 - Facture 123" in failed[6]
    assert "Erreur : Date invalide" in failed[6]
    assert "Reçu : 31/13/2026" in failed[6]
    assert get_rows(1, 2)[0][2] == "bbgr.csv"


def test_export_excel(traces):
    file_io = io.BytesIO()

    # Nom construit comme dans la vue : LECTURES_TRACES_{Y_M_D}{timestamp}.xlsx
    result = excel_edi_traces(file_io, "LECTURES_TRACES_2026_10_121760000000.xlsx", 0, 50)

    file_io.seek(0)
    sheet = load_workbook(file_io, read_only=True).active
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    assert "OK" in result
    assert sheet.title == "LECTURES TRACES"
    assert any(row[2] == "bbgr.csv" and row[4] == "X" for row in rows)


def test_vue_export(logged_client, monkeypatch, traces):
    """GET renvoie le fichier, POST ou erreur redirigent vers l'écran des traces"""
    url = reverse("traces:edi_traces_export", kwargs={"start_index": 0, "end_index": 50})

    response = logged_client.get(url)
    assert response.status_code == 200
    assert response["Content-Type"] == CONTENT_TYPE_EXCEL
    assert "LECTURES_TRACES_" in response["Content-Disposition"]

    posted = logged_client.post(url)
    assert posted.status_code == 302
    assert posted.url == reverse("traces:edi_traces")

    def _boom(*_args, **_kwargs):
        raise RuntimeError("export cassé")

    monkeypatch.setattr(traces_views, "response_file", _boom)
    failed = logged_client.get(url)
    assert failed.status_code == 302
    assert failed.url == reverse("traces:edi_traces")
