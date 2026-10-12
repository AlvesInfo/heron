# pylint: disable=E0401,W0621
"""FR : Tests de la configuration, des modèles, des formulaires et des exports Excel de l'app rfa


EN : Tests for rfa app configuration, models, forms and Excel exports

created at: 2026-10-12
created by: Paulo ALVES
"""

import importlib
import io
from decimal import Decimal

import pytest
from django.apps import apps as django_apps
from django.db import IntegrityError, transaction
from openpyxl import load_workbook

from apps.rfa import urls as rfa_urls
from apps.rfa.apps import RfaConfig
from apps.rfa.excel_outputs import (
    rfa_clients_exclusion,
    rfa_clients_suppliers_exclusion,
    rfa_sections,
    rfa_sections_pro_exclusions,
    rfa_signboards_exclusion,
    rfa_suppliers_rate,
)
from apps.rfa.forms import (
    AxeRfaForm,
    ClientExclusionForm,
    ClientSupplierExclusionForm,
    DeleteClientExclusionForm,
    DeleteClientSupplierExclusionForm,
    DeleteSectionProExclusionForm,
    DeleteSectionRfaForm,
    DeleteSignboardExclusionForm,
    DeleteSupplierRateForm,
    SectionProExclusionForm,
    SectionRfaForm,
    SignboardExclusionForm,
    SupplierRateForm,
)
from apps.rfa.models import SignboardExclusion, SupplierRate
from apps.rfa.tests.conftest import RFA_PERIOD

pytestmark = pytest.mark.django_db


def _sheet_rows(file_io):
    file_io.seek(0)
    return [
        tuple(row)
        for row in load_workbook(file_io, read_only=True).active.iter_rows(values_only=True)
    ]


def test_app_config_et_routes():
    config = django_apps.get_app_config("rfa")
    names = [pattern.name for pattern in rfa_urls.urlpatterns]

    assert isinstance(config, RfaConfig)
    assert config.name == "apps.rfa"
    assert len(list(config.get_models())) == 6
    assert len(names) == len(set(names)) == 31
    assert importlib.import_module("apps.rfa.admin") is not None


def test_unicite_des_parametrages(make_signboard_exclusion, make_supplier_rate):
    exclusion = make_signboard_exclusion("ACF")
    rate = make_supplier_rate("FOUR001", "0.05")

    assert exclusion.signboard.code == "ACF"
    assert rate.rfa_rate == Decimal("0.05")

    with pytest.raises(IntegrityError), transaction.atomic():
        SignboardExclusion.objects.create(signboard=exclusion.signboard)

    with pytest.raises(IntegrityError), transaction.atomic():
        SupplierRate.objects.create(supplier=rate.supplier, rfa_rate=1)


# FORMULAIRES ======================================================================================


def test_formulaires_de_saisie(make_signboard, make_cct, make_maison, make_society, make_section):
    """Seules les maisons de type vente dont le cct est actif sont proposées à l'exclusion"""
    make_cct("AF0001", active=True)
    make_cct("AF0002", active=True)
    maison = make_maison("AF0001", type_x3=1)
    exclu = make_maison("AF0002", type_x3=0)
    make_maison("AF0004", type_x3=1)
    supplier = make_society("FOUR001", is_client=False)
    section_rfa = make_section("RFA", RFA_PERIOD)
    section_pro = make_section("PRO", "OPT")
    make_signboard("ACF")

    assert SignboardExclusionForm({"signboard": "ACF"}).is_valid()
    assert SignboardExclusionForm({"signboard": "XXX"}).is_valid() is False
    assert ClientExclusionForm({"cct": "AF0001"}).is_valid()
    assert ClientExclusionForm({"cct": "AF0002"}).is_valid() is False
    assert ClientExclusionForm({"cct": "AF0004"}).is_valid() is False
    assert exclu.type_x3.num == 0
    assert ClientSupplierExclusionForm(
        {"maison": maison.uuid_identification, "supplier": "FOUR001"}
    ).is_valid()
    assert SupplierRateForm({"supplier": "FOUR001", "rfa_rate": "0.05"}).is_valid()
    assert (
        SupplierRateForm({"supplier": supplier.third_party_num, "rfa_rate": ""}).is_valid() is False
    )
    assert SectionRfaForm({"axe_rfa": section_rfa.uuid_identification}).is_valid()
    assert SectionRfaForm({"axe_rfa": section_pro.uuid_identification}).is_valid() is False
    assert SectionProExclusionForm({"axe_pro": section_pro.uuid_identification}).is_valid()
    assert SectionProExclusionForm({"axe_pro": section_rfa.uuid_identification}).is_valid() is False


def test_formulaires_de_suppression(
    make_signboard_exclusion,
    make_client_exclusion,
    make_client_supplier_exclusion,
    make_supplier_rate,
    make_section_rfa,
    make_section_pro_exclusion,
):
    cases = [
        (DeleteSignboardExclusionForm, make_signboard_exclusion()),
        (DeleteClientExclusionForm, make_client_exclusion()),
        (DeleteClientSupplierExclusionForm, make_client_supplier_exclusion()),
        (DeleteSupplierRateForm, make_supplier_rate()),
        (DeleteSectionRfaForm, make_section_rfa()),
        (DeleteSectionProExclusionForm, make_section_pro_exclusion()),
    ]

    for form_class, instance in cases:
        assert form_class({"id": instance.pk}).is_valid()
        assert form_class({"id": instance.pk + 1000}).is_valid() is False
        assert form_class({"id": ""}).is_valid() is False


def test_formulaire_axe_rfa(make_section, make_edi_validation):
    """Les sections de l'axe RFA sont proposées triées, la période en cours est présélectionnée"""
    make_section("RFA", "2026RFA_02")
    make_section("RFA", RFA_PERIOD)
    make_section("PRO", "OPT")
    make_edi_validation()

    form = AxeRfaForm()
    valid = AxeRfaForm({"section": RFA_PERIOD})

    assert form.fields["section"].choices == [
        (RFA_PERIOD, RFA_PERIOD),
        ("2026RFA_02", "2026RFA_02"),
    ]
    assert form.fields["section"].initial == RFA_PERIOD
    assert valid.is_valid()
    assert valid.cleaned_data["section"] == RFA_PERIOD
    assert AxeRfaForm({"section": "OPT"}).is_valid() is False


# EXPORTS EXCEL ====================================================================================


def test_export_enseignes(make_signboard_exclusion):
    make_signboard_exclusion("ACF")
    file_io = io.BytesIO()

    result = rfa_signboards_exclusion.excel_list_rfa_signboards_exclusion(
        file_io, "LISTING_DES_EXCLUSIONS_RFA_ENSEIGNES_2026_10_12_1.xlsx"
    )

    assert "OK" in result
    assert list(rfa_signboards_exclusion.get_row()) == [("ACF", "Enseigne ACF")]
    assert ("ACF", "Enseigne ACF") in _sheet_rows(file_io)


def test_export_clients(make_client_exclusion, make_maison):
    make_client_exclusion(make_maison("AF0001", type_x3=1))
    file_io = io.BytesIO()

    result = rfa_clients_exclusion.excel_list_rfa_clients_exclusion(
        file_io, "LISTING_DES_EXCLUSIONS_RFA_CLIENTS_CCT_2026_10_12_1.xlsx"
    )

    assert "OK" in result
    assert list(rfa_clients_exclusion.get_row()) == [("AF0001", "Maison AF0001")]
    assert ("AF0001", "Maison AF0001") in _sheet_rows(file_io)


def test_export_clients_fournisseurs(make_client_supplier_exclusion, make_maison):
    make_client_supplier_exclusion(make_maison("AF0001", type_x3=1), "FOUR001")
    file_io = io.BytesIO()

    result = rfa_clients_suppliers_exclusion.excel_list_rfa_clients_suppliers_exclusion(
        file_io, "LISTING_DES_EXCLUSIONS_RFA_CLIENTS_CCT_2026_10_12_1.xlsx"
    )

    assert "OK" in result
    assert list(rfa_clients_suppliers_exclusion.get_row()) == [
        ("AF0001", "Maison AF0001", "FOUR001")
    ]
    assert ("AF0001", "Maison AF0001", "FOUR001") in _sheet_rows(file_io)


def test_export_taux(make_supplier_rate):
    make_supplier_rate("FOUR001", "0.05", name="Fournisseur Un")
    file_io = io.BytesIO()

    result = rfa_suppliers_rate.excel_list_suppliers_rate(
        file_io, "LISTING_DES_TAUX_RFA_2026_10_12_1.xlsx"
    )

    assert "OK" in result
    assert list(rfa_suppliers_rate.get_row()) == [("FOUR001", "Fournisseur Un", Decimal("0.05"))]
    assert ("FOUR001", "Fournisseur Un", 0.05) in _sheet_rows(file_io)


def test_export_sections(make_section_rfa, make_section_pro_exclusion):
    make_section_rfa(RFA_PERIOD)
    make_section_pro_exclusion("OPT")
    rfa_io, pro_io = io.BytesIO(), io.BytesIO()

    rfa_result = rfa_sections.excel_list_rfa_sections(
        rfa_io, "LISTING_DES_AXES_RFA_2026_10_12_1.xlsx"
    )
    pro_result = rfa_sections_pro_exclusions.excel_list_rfa_sections_pro_exclusions(
        pro_io, "LISTING_DES_EXCLUSIONS_RFA_AXE_PRO_2026_10_12_1.xlsx"
    )

    assert "OK" in rfa_result
    assert "OK" in pro_result
    assert list(rfa_sections.get_row()) == [(RFA_PERIOD, f"Section {RFA_PERIOD}")]
    assert list(rfa_sections_pro_exclusions.get_row()) == [("OPT", "Section OPT")]
    assert (RFA_PERIOD, f"Section {RFA_PERIOD}") in _sheet_rows(rfa_io)
    assert ("OPT", "Section OPT") in _sheet_rows(pro_io)
